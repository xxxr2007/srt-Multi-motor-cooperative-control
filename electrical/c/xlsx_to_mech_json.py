#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsx_to_mech_json.py —— 机械组手填 Excel → 仿真参数 json 的转换器
====================================================================
背景
    机械组同学用 Excel 填参数最顺手，但 git 对二进制 xlsx 无法 diff、无法合并，
    所以仓库的"真源"仍是 data/mechanical/mech_params.json（文本格式）。
    本脚本就是两者之间的桥：Excel 负责"人写舒服"，json 负责"机器读准确"。

两种用法（都在仓库根目录执行；无 openpyxl 时用 uv 临时环境，见文件头命令）
    ① 生成手填模板（首次使用 / 模板丢了）：
        python electrical/c/xlsx_to_mech_json.py --make-template
        -> 生成 mechanical/params/motor_params.xlsx（带说明列，机械组照着填）

    ② 填完后转成真源（机械组每次改完参数跑一次）：
        python electrical/c/xlsx_to_mech_json.py
        -> 覆盖写 data/mechanical/mech_params.json（之后电控组跑 merge_params.py 生效）

    没有 openpyxl 时的运行方式（推荐写进组内文档）：
        uv run --no-project --python 3.12 --with openpyxl python electrical/c/xlsx_to_mech_json.py

数据流向（与 docs/collaboration.md 一致）
    mechanical/params/motor_params.xlsx   机械组手填
        ↓ 本脚本（②模式）
    data/mechanical/mech_params.json      参数真源（进 git，可追溯）
        ↓ electrical/c/merge_params.py
    data/params.json + build/params_generated.h   合并产物（自动生成，勿手改）
"""

import argparse            # 命令行参数：区分 --make-template 与转换模式
import datetime            # 自动填今天的日期（ISO 格式），省得手写错格式
import json                # 写 mech_params.json（ensure_ascii=False 保住中文可读性）
import os                  # 路径拼接：脚本挪到 electrical/c/ 后按自身位置找仓库根
import sys                 # 出错时带退出码终止，方便批处理/CI 中断

# 第三方库 openpyxl：Excel 读写。标准库不带，需 pip/uv 安装（见文件头命令）。
try:
    import openpyxl        # .xlsx 读写库（模板生成 + 解析共用）
except ImportError:        # 给出人话提示 + 可直接复制的安装命令，而不是一串 traceback
    print("[ERR] 缺少 openpyxl 库。任选一种安装方式：")
    print("      pip install openpyxl")
    print('      uv run --no-project --with openpyxl python electrical/c/xlsx_to_mech_json.py')
    sys.exit(2)

# ---------------------------------------------------------------- 路径解析
HERE = os.path.dirname(os.path.abspath(__file__))          # .../仓库根/electrical/c
ROOT = os.path.dirname(os.path.dirname(HERE))              # 上两级 = 仓库根（c → electrical → 根）
XLSX = os.path.join(ROOT, "mechanical", "params", "motor_params.xlsx")   # 机械组手填表
MECH_JSON = os.path.join(ROOT, "data", "mechanical", "mech_params.json")  # 参数真源（转换目标）

# ---------------------------------------------------------------- 模板内容定义
# 每行：(参数名, 默认值, 单位, 填写说明)。集中定义一份，生成模板与解析共用同一张"字段表"，
# 以后加字段只需在这里补一行，模板和转换逻辑自动同步——避免两处维护对不上。
# "待填"字符串是占位标记：转换时若值仍是它，说明机械组没填，会报警告。
FIELDS = [
    # ---- 版本与来源（meta）----
    ("version",       "v1",                 "-",   "每交付一版 +1（v1, v2...），供 git 对账"),
    ("delivered_by",  "待填：姓名",          "-",   "谁填的表，方便追责问询"),
    ("source",        "待填：参数来源",      "-",   "CATIA 测量惯量 / 联轴器手册 / 实验辨识，写清来源"),

    # ---- 电机参数（motor）----
    ("model",         "待填：电机型号",      "-",   "产品型号，例：某 400W 直流伺服电机"),
    ("J1",            0.010,                "kg*m^2", "1号机转子+折算负载总惯量（绕电机轴）"),
    ("J2",            0.011,                "kg*m^2", "2号机（±10% 摄动体现多电机不一致）"),
    ("J3",            0.0095,               "kg*m^2", "3号机"),
    ("J4",            0.0105,               "kg*m^2", "4号机"),
    ("J_nominal",     0.01025,              "kg*m^2", "四台的名义均值（观测器用 J_n = 这个值）"),
    ("B1",            0.0012,               "N*m*s/rad", "1号机粘性摩擦系数（手册/辨识，CATIA 算不出）"),
    ("B2",            0.0011,               "N*m*s/rad", "2号机"),
    ("B3",            0.0013,               "N*m*s/rad", "3号机"),
    ("B4",            0.00115,              "N*m*s/rad", "4号机"),
    ("B_source",      "待填：B 的来源",      "-",   "手册值 / 经验值 / 实验辨识，三选一写明"),
    ("mass_kg",       None,                 "kg",   "电机本体质量（选填，报告用）"),
    ("perturb_pct",   10,                   "%",    "参数摄动幅度（现在 ±10%，可按公差改）"),

    # ---- 联轴器与传动（coupling）----
    ("scheme",        "待填：刚性/弹性柱销/膜片", "-", "联轴器类型（选型结果）"),
    ("Ks",            None,                 "N*m/rad", "扭转刚度 —— 查联轴器手册，机械组核心交付"),
    ("Ds",            None,                 "N*m*s/rad", "阻尼系数（手册给出就有，没有填 0）"),
    ("J_motor_side",  0.010,                "kg*m^2", "双惯量模型电机侧惯量 J1（弹性方案 A 用）"),
    ("J_load_side",   0.010,                "kg*m^2", "双惯量模型负载侧惯量 J2"),
    ("omega_n",       None,                 "rad/s", "机械谐振频率 = sqrt(Ks*(1/J1+1/J2))，必算"),
    ("i_ratio",       1.0,                  "-",    "减速比 i（负载侧到电机侧折算：J/i^2）"),
    ("eta",           1.0,                  "-",    "传动效率（折算负载转矩用：T_L/(i*eta)）"),

    # ---- 负载工况物理来源（load_physics）----
    ("load_step",     "待填：突加负载的工况",  "-",   "例：传送带突然上料 / 单机卡滞"),
    ("load_sin",      "待填：波动负载的工况",  "-",   "例：旋转刀盘周期切削"),
    ("load_ramp",     "待填：斜坡负载的工况",  "-",   "例：逐渐加压的压紧过程"),
]

# 单元格表头（模板第一行）；列宽手动调过，保证说明列能读全
HEADERS = ["参数名", "值（机械组填这列）", "单位", "填写说明"]


def make_template():
    """生成 mechanical/params/motor_params.xlsx 手填模板。

    模板特征：参数名+单位+说明由脚本写入（只读参考），机械组只改"值"那一列；
    转换时也只认"值"列，这样填写界面和数据结构永远一致。
    """
    wb = openpyxl.Workbook()                    # 新建工作簿（默认含一张空 sheet）
    ws = wb.active                              # 取默认 sheet
    ws.title = "电机与传动参数"                   # 表名即内容，打开就知道填什么

    ws.append(HEADERS)                          # 第一行：表头
    for name, default, unit, desc in FIELDS:    # 每个字段一行：默认值带进去当参考
        ws.append([name, default, unit, desc])

    # ---- 版式美化：不影响数据，只是让机械组填得舒服 ----
    ws.column_dimensions["A"].width = 16         # 参数名列
    ws.column_dimensions["B"].width = 22         # 值列（要手填，给宽一点）
    ws.column_dimensions["C"].width = 14         # 单位列
    ws.column_dimensions["D"].width = 46         # 说明列（最长，占大头）
    ws.freeze_panes = "A2"                       # 冻结表头行：往下滚时表头不消失

    os.makedirs(os.path.dirname(XLSX), exist_ok=True)   # 目录不存在则建（机械/params/）
    wb.save(XLSX)                               # 落盘
    print(f"[OK] 模板已生成：{os.path.relpath(XLSX, ROOT)}")
    print("     机械组只填 B 列（值），填完运行本脚本（不带参数）转换为 json。")


def parse_template():
    """读 xlsx 的"值"列 → 组装成 mech_params.json 的嵌套结构。

    映射关系（与 json 的键一一对应，改哪加哪都要两边同步）：
        version/delivered_by/source -> meta
        model/J*/B*/...            -> motor（J1..J4、B1..B4 收拢成数组）
        scheme/Ks/...              -> coupling
        load_step/sin/ramp         -> load_physics
    """
    if not os.path.isfile(XLSX):                # 表还没生成/没放进仓库就先提示
        print(f"[ERR] 找不到 {os.path.relpath(XLSX, ROOT)}，先运行 --make-template 生成模板")
        sys.exit(2)

    wb = openpyxl.load_workbook(XLSX)           # 只读数据，不改表
    ws = wb.active

    raw = {}                                    # {参数名: 值} 的扁平字典，先收进来
    for row in ws.iter_rows(min_row=2, values_only=True):   # 跳过表头，从数据行开始
        if row is None or row[0] is None:
            continue                            # 空行直接跳过（Excel 常见尾部空行）
        name, value = str(row[0]).strip(), row[1]
        raw[name] = value                       # 只取"值"列；单位/说明列不进数据

    # ---- 完整性检查：模板里定义过的字段必须都出现在表里 ----
    missing = [f[0] for f in FIELDS if f[0] not in raw]
    if missing:
        print(f"[ERR] 表里缺字段：{missing} —— 请勿删行；模板坏了就重跑 --make-template")
        sys.exit(2)

    # ---- 占位值检查：还留着"待填"开头的值 = 机械组没填完，逐个点名 ----
    unfilled = [k for k, v in raw.items()
                if isinstance(v, str) and v.startswith("待填")]
    if unfilled:
        print(f"[WARN] 以下字段仍是占位值（未填）：{unfilled}")
        print("       照常转换，但这些值到合并校验时大概率过不了量级检查。")

    def num(key):
        """取数值字段并转 float；空/非数直接报错终止（宁可硬失败也不带错值进仿真）。"""
        v = raw[key]
        if v is None or (isinstance(v, str) and not v.strip()):
            print(f"[ERR] 数值字段 {key} 为空 —— 必填（不知道就先填手册典型值并注明来源）")
            sys.exit(2)
        try:
            return float(v)                     # Excel 数值单元格进来就是 int/float
        except (TypeError, ValueError):
            print(f"[ERR] 字段 {key} 的值 [{v}] 不是数字")
            sys.exit(2)

    def text(key, default=""):
        """取文本字段；空值退回默认字符串而不是 None（json 里保持类型整齐）。"""
        v = raw[key]
        return str(v).strip() if v is not None else default

    # ---- 按 mech_params.json 的嵌套结构组装 ----
    out = {
        "meta": {                               # 版本信息：git 历史之外的"表内版本"，双保险
            "version": text("version", "v1"),
            "delivered_by": text("delivered_by"),
            "date": datetime.date.today().isoformat(),   # 转换当天日期，自动生成
            "source": text("source"),
            "note": "本文件由 mechanical/params/motor_params.xlsx 经 xlsx_to_mech_json.py 转换生成",
        },
        "motor": {
            "model": text("model"),
            "J": [num(f"J{i}") for i in range(1, 5)],     # J1..J4 收拢成数组（仿真按台索引）
            "J_nominal": num("J_nominal"),
            "B": [num(f"B{i}") for i in range(1, 5)],     # B1..B4 同上
            "B_source": text("B_source"),
            "mass_kg": num("mass_kg") if raw.get("mass_kg") not in (None, "") else None,
            "perturb_pct": num("perturb_pct"),
            "perturb_reasons": ["制造公差", "温升", "磨损", "装配差异"],  # 依据固定四项，报告里展开
        },
        "coupling": {
            "scheme": text("scheme"),
            "Ks": num("Ks") if raw.get("Ks") not in (None, "") else None,      # 没选型完允许空
            "Ds": num("Ds") if raw.get("Ds") not in (None, "") else None,
            "J1": num("J_motor_side"),          # json 键名 J1/J2 = 双惯量模型两侧
            "J2": num("J_load_side"),
            "omega_n_rad_s": num("omega_n") if raw.get("omega_n") not in (None, "") else None,
            "I_ratio": num("i_ratio"),
            "eta": num("eta"),
            "note": "Ks 查联轴器手册；ω_n 由机械组算好后交给电控组，观测器带宽需覆盖其 2~5 倍",
        },
        "load_physics": {                       # 三种负载各自的"物理是什么"（评审必问）
            "step": text("load_step"),
            "sin": text("load_sin"),
            "ramp": text("load_ramp"),
        },
    }

    # ---- 写出真源 json（缩进 2 格保可读；ensure_ascii=False 中文不转义）----
    os.makedirs(os.path.dirname(MECH_JSON), exist_ok=True)
    with open(MECH_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
        f.write("\n")                           # 结尾补换行，符合 POSIX 文本文件惯例

    print(f"[OK] 已转换：{os.path.relpath(XLSX, ROOT)} -> {os.path.relpath(MECH_JSON, ROOT)}")
    print("     下一步：电控组跑 electrical/c/merge_params.py 合并参数，再重新编译。")


def main():
    """入口：--make-template 生成模板；默认转换 xlsx -> json。"""
    ap = argparse.ArgumentParser(description="机械组 Excel 参数表 <-> 仿真参数 json 转换器")
    ap.add_argument("--make-template", action="store_true",
                    help="生成 mechanical/params/motor_params.xlsx 手填模板（首次使用）")
    args = ap.parse_args()

    if args.make_template:
        make_template()                         # 模式①：出一张空表给机械组
    else:
        parse_template()                        # 模式②：表 -> json 真源


if __name__ == "__main__":
    main()
